"""
TemplateFill — HTML 转 python-docx 转换器

将 QTextEdit 输出的 HTML 转换为 python-docx 段落和格式化 runs，
插入到文档中指定位置之后。

支持的内联格式：b/strong, i/em, u, s/strike/del, span[style]
支持的块级元素：p, h1-h6, ul/ol/li, br, table（真实表格）
动态章节块类型：heading / text / table / image
"""

from lxml import html as lxml_html
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Pt, Cm, Emu, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ── 内部常量 ────────────────────────────────────────────

BODY_MARKER = 'BODY_CONTENT_MARKER_9F8E7D6C5B4A_PLACEHOLDER'

# ── 主入口 ──────────────────────────────────────────────

def html_to_docx(html_string, doc, ref_element=None):
    """
    将 HTML 字符串转换为 python-docx 段落，插入到 ref_element 之后。

    参数:
        html_string: HTML 文本（来自 QTextEdit.toHtml()）
        doc: python-docx / docxtpl Document 对象
        ref_element: lxml <w:p> 元素，新内容插入到此之后。
                     若为 None，追加到文档末尾。

    返回: 最后插入的 XML 元素（用于链式调用）
    """
    if not html_string or not html_string.strip():
        return ref_element

    # 判断是否为完整 HTML 文档
    stripped = html_string.strip().lower()
    if stripped.startswith('<!doctype') or stripped.startswith('<html'):
        # 完整文档 — 提取 body
        tree = lxml_html.fromstring(html_string)
        body = tree.find('.//body')
        root = body if body is not None else tree
    else:
        # HTML 片段 — 包装在 div 中确保多根元素被正确处理
        root = lxml_html.fromstring(f'<div>{html_string}</div>')

    current_ref = ref_element

    for child in root:
        if not isinstance(child.tag, str):
            continue
        current_ref = _process_block(child, doc, current_ref)

    # 如果根元素有直接文本但无子元素，作为段落处理
    if root.text and root.text.strip() and not list(root):
        current_ref = _process_paragraph(root, doc, current_ref)

    return current_ref


# ── 块级元素处理 ─────────────────────────────────────────

def _process_block(node, doc, ref):
    """处理块级 HTML 元素，返回最后插入的 XML 元素"""
    tag = node.tag.lower()

    if tag in ('p', 'div', 'span'):
        # 段落 — 可能含内联格式
        return _process_paragraph(node, doc, ref)
    elif tag in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6'):
        level = int(tag[1])
        return _process_heading(node, doc, ref, level)
    elif tag == 'ul':
        return _process_list(node, doc, ref, ordered=False)
    elif tag == 'ol':
        return _process_list(node, doc, ref, ordered=True)
    elif tag == 'br':
        # 顶层 br — 插入空段落
        return _insert_paragraph(doc, ref)[1]
    elif tag in ('table',):
        # 表格 — 转换为真实 Word 表格
        return _process_table_html(node, doc, ref)
    elif tag in ('meta', 'style', 'head', 'title', 'link'):
        # 忽略 HTML 头部元素
        return ref
    else:
        # 未知块级元素 — 处理子节点
        if node.text and node.text.strip():
            return _process_paragraph(node, doc, ref)
        current_ref = ref
        for child in node:
            if isinstance(child.tag, str):
                current_ref = _process_block(child, doc, current_ref)
        return current_ref


def _process_paragraph(node, doc, ref):
    """处理 <p> 元素"""
    new_p, new_ref = _insert_paragraph(doc, ref)
    _process_inline(node, new_p, {})
    return new_ref


def _process_heading(node, doc, ref, level):
    """处理 <h1>-<h6> 标题"""
    new_p, new_ref = _insert_paragraph(doc, ref)
    # 尝试使用 Word 内置标题样式
    style_name = f'Heading {level}'
    try:
        new_p.style = doc.styles[style_name]
    except (KeyError, IndexError):
        # 样式不存在 — 手动设置格式
        _apply_manual_heading_format(new_p, level)
    _process_inline(node, new_p, {'bold': True})
    return new_ref


def _process_list(node, doc, ref, ordered=False):
    """处理 <ul>/<ol> 列表"""
    current_ref = ref
    counter = 1
    for child in node:
        if not isinstance(child.tag, str):
            continue
        if child.tag.lower() == 'li':
            new_p, new_ref = _insert_paragraph(doc, current_ref)
            # 尝试使用列表样式
            style_name = 'List Number' if ordered else 'List Bullet'
            try:
                new_p.style = doc.styles[style_name]
            except (KeyError, IndexError):
                # 样式不存在 — 手动添加前缀
                prefix = f'{counter}. ' if ordered else '\u2022 '
                run = new_p.add_run(prefix)
                run.bold = True
            _process_inline(child, new_p, {})
            current_ref = new_ref
            counter += 1
    return current_ref


def _process_table_html(node, doc, ref):
    """处理 HTML <table> — 提取行列文本后生成真实 Word 表格"""
    rows_data = []
    has_header = False
    # 支持 thead/tbody/tfoot 包裹与裸 tr
    for tr in node.iter('tr'):
        if not isinstance(tr.tag, str):
            continue
        row = []
        for td in tr:
            if isinstance(td.tag, str) and td.tag.lower() in ('td', 'th'):
                row.append(td.text_content().strip())
                if td.tag.lower() == 'th':
                    has_header = True
        if row:
            rows_data.append(row)
    if not rows_data:
        return ref
    tbl_el, tail = _insert_table(doc, ref, rows_data, header=has_header)
    return tail


# ── 表格插入 ─────────────────────────────────────────────

def _insert_table(doc, ref, rows_data, header=False):
    """
    在 ref 之后插入表格（单元格为纯文本）。

    参数:
        rows_data: [[c00, c01, ...], [c10, ...], ...] 字符串二维数组
                   （行长度不齐时自动右侧补空）
        header:    首行是否为表头（加粗 + 浅色底纹）

    返回: (tbl XML 元素, 表格后的空段落元素) — 后者用作新的 ref；
          无有效数据时返回 (ref, ref)
    """
    n_rows = len(rows_data)
    n_cols = max((len(r) for r in rows_data), default=0)
    if n_rows == 0 or n_cols == 0:
        return ref, ref

    # 补齐不整齐的行，统一转字符串
    norm = []
    for row in rows_data:
        row = list(row) + [''] * (n_cols - len(row))
        norm.append([str(c) if c is not None else '' for c in row])

    # 先在文档末尾创建表格，再移动到目标位置
    table = doc.add_table(rows=n_rows, cols=n_cols)
    try:
        table.style = doc.styles['Table Grid']
    except (KeyError, IndexError):
        # 模板没有 Table Grid 样式 — 手动加边框兜底
        _apply_table_borders(table)
    table.autofit = True

    for r, row in enumerate(norm):
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            para = cell.paragraphs[0]
            run = para.add_run(text)
            run.font.size = Pt(10.5)
            if header and r == 0:
                run.bold = True
                _shade_cell(cell, 'DCE6F1')

    # 摘出表格 XML，插到 ref 之后
    tbl_el = table._tbl
    tbl_el.getparent().remove(tbl_el)
    if ref is not None:
        ref.addnext(tbl_el)
    else:
        doc.element.body.append(tbl_el)

    # 表格后补一个空段落（Word 规范：表格不能直接跟表格/结尾）
    tail_p = OxmlElement('w:p')
    tbl_el.addnext(tail_p)
    return tbl_el, tail_p


def _apply_table_borders(table):
    """手动给表格加全边框（模板缺 Table Grid 样式时的兜底）"""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:color'), '999999')
        borders.append(el)
    tbl_pr.append(borders)


def _shade_cell(cell, fill_hex):
    """单元格底纹（如表头浅蓝 DCE6F1）"""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), fill_hex)
    tc_pr.append(shd)


# ── 图片插入 ─────────────────────────────────────────────

def _insert_image(doc, ref, image_path, width_cm=0, caption=''):
    """
    在 ref 之后插入图片（居中）+ 可选图注段落。

    参数:
        image_path: 本地图片文件路径
        width_cm:   显示宽度（厘米）。0 = 原始大小（最宽 15 cm）
        caption:    图注文本（可选，居中、9pt、灰色）

    返回: 最后插入的 XML 元素（新的 ref）
    """
    new_p, new_ref = _insert_paragraph(doc, ref)
    new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    try:
        run = new_p.add_run()
        pic = run.add_picture(image_path)
    except Exception:
        # 文件不存在或格式不支持 — 插入占位提示
        ph = new_p.add_run(f'[图片缺失或格式不支持：{image_path}]')
        ph.italic = True
        ph.font.color.rgb = RGBColor(0xE7, 0x3C, 0x3C)
        return new_ref

    # 宽度控制（指定宽度时高度按比例缩放）
    try:
        target = Cm(float(width_cm)) if width_cm else None
    except (TypeError, ValueError):
        target = None

    if target is not None:
        if pic.width and pic.width != target:
            ratio = target / pic.width
            pic.height = Emu(int(pic.height * ratio))
        pic.width = target
    else:
        # 原始大小，最宽 15 cm（A4 页宽安全值）
        max_w = Cm(15)
        if pic.width and pic.width > max_w:
            ratio = max_w / pic.width
            pic.height = Emu(int(pic.height * ratio))
            pic.width = max_w

    last_ref = new_ref

    # 图注段落
    if caption and str(caption).strip():
        cap_p, cap_ref = _insert_paragraph(doc, new_ref)
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = cap_p.add_run(str(caption).strip())
        cap_run.font.size = Pt(9)
        cap_run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        last_ref = cap_ref

    return last_ref


# ── 内联格式处理 ─────────────────────────────────────────

def _process_inline(node, paragraph, fmt):
    """
    递归处理内联内容，将文本和格式写入 paragraph。

    fmt: 当前格式状态 dict {
        bold: bool, italic: bool, underline: bool,
        strikethrough: bool
    }
    """
    # 1. 处理节点前文本
    if node.text:
        _add_run(paragraph, node.text, fmt)

    # 2. 处理子元素
    for child in node:
        if not isinstance(child.tag, str):
            continue
        ctag = child.tag.lower()

        if ctag == 'br':
            run = paragraph.add_run()
            run.add_break()
        elif ctag in ('b', 'strong'):
            _process_inline(child, paragraph, {**fmt, 'bold': True})
        elif ctag in ('i', 'em'):
            _process_inline(child, paragraph, {**fmt, 'italic': True})
        elif ctag == 'u':
            _process_inline(child, paragraph, {**fmt, 'underline': True})
        elif ctag in ('s', 'strike', 'del'):
            _process_inline(child, paragraph, {**fmt, 'strikethrough': True})
        elif ctag == 'span':
            # QTextEdit 用 span + style 代替 b/i/u
            child_fmt = _parse_span_style(child, fmt)
            _process_inline(child, paragraph, child_fmt)
        elif ctag == 'a':
            # 超链接 — 当作普通文本
            _process_inline(child, paragraph, fmt)
        elif ctag in ('p', 'div', 'h1', 'h2', 'h3', 'h4', 'h5', 'h6',
                       'ul', 'ol', 'li', 'table', 'tr', 'td', 'th',
                       'tbody', 'thead', 'tfoot'):
            # 块级元素出现在内联上下文中 — 递归处理其内联内容
            _process_inline(child, paragraph, fmt)
        else:
            # 未知内联元素 — 递归
            _process_inline(child, paragraph, fmt)

        # 3. 处理子元素的 tail 文本（属于当前 fmt 上下文）
        if child.tail:
            _add_run(paragraph, child.tail, fmt)


def _parse_span_style(node, parent_fmt):
    """解析 <span style="..."> 的 CSS 样式，合并到格式 dict"""
    fmt = dict(parent_fmt)
    style = node.get('style', '')
    if not style:
        return fmt

    style_lower = style.lower()

    # font-weight
    if 'font-weight' in style_lower:
        weight_part = _extract_css_value(style_lower, 'font-weight')
        if weight_part in ('bold', '600', '700', '800', '900'):
            fmt['bold'] = True
        elif weight_part in ('normal', '400', 'lighter', '300', '200', '100'):
            fmt['bold'] = False

    # font-style
    if 'font-style' in style_lower:
        if 'italic' in _extract_css_value(style_lower, 'font-style'):
            fmt['italic'] = True

    # text-decoration
    if 'text-decoration' in style_lower:
        td = _extract_css_value(style_lower, 'text-decoration')
        if 'underline' in td:
            fmt['underline'] = True
        if 'line-through' in td:
            fmt['strikethrough'] = True

    return fmt


def _extract_css_value(style_str, prop):
    """从 CSS 字符串中提取指定属性的值"""
    for part in style_str.split(';'):
        part = part.strip()
        if part.startswith(prop):
            colon_idx = part.find(':')
            if colon_idx >= 0:
                return part[colon_idx + 1:].strip()
    return ''


def _add_run(paragraph, text, fmt):
    """添加带格式的 run"""
    if not text:
        return
    run = paragraph.add_run(text)
    if fmt.get('bold'):
        run.bold = True
    if fmt.get('italic'):
        run.italic = True
    if fmt.get('underline'):
        run.underline = True
    if fmt.get('strikethrough'):
        run.font.strike = True


# ── 段落插入辅助 ─────────────────────────────────────────

def _insert_paragraph(doc, ref_element):
    """
    在 ref_element 之后插入新段落。
    返回 (Paragraph 对象, 新的 XML 元素)
    """
    new_p = OxmlElement('w:p')
    if ref_element is not None:
        ref_element.addnext(new_p)
    else:
        # 追加到文档 body 末尾
        doc.element.body.append(new_p)
    para = Paragraph(new_p, doc)
    return para, new_p


def _apply_manual_heading_format(paragraph, level):
    """当标题样式不存在时，手动设置格式"""
    sizes = {1: 18, 2: 16, 3: 14, 4: 13, 5: 12, 6: 11}
    for run in paragraph.runs:
        run.bold = True
        run.font.size = Pt(sizes.get(level, 12))


# ── 动态章节插入 ─────────────────────────────────────────

def find_body_marker(doc):
    """
    在已渲染的文档中查找 body marker 段落。
    返回 Paragraph 对象，未找到返回 None。
    """
    for para in doc.paragraphs:
        if BODY_MARKER in para.text:
            return para
    return None


def insert_sections(doc, sections):
    """
    在 body marker 位置插入动态章节内容。

    参数:
        doc: 已渲染的 DocxTemplate 对象
        sections: list of {
            'title': str,           # 章节标题（H1）
            'blocks': list of {     # 子项
                'type': 'heading' | 'text' | 'table' | 'image',
                'text': str,         # heading: 标题文本
                'level': int,        # heading: 级别 (2-6)
                'html': str,         # text: HTML 内容
                'data': [[...]],     # table: 字符串二维数组
                'header': bool,      # table: 首行是否为表头
                'path': str,         # image: 本地图片路径
                'width': float,      # image: 显示宽度 cm（0=自动）
                'caption': str,      # image: 图注（可选）
            }
        }
    """
    marker_para = find_body_marker(doc)
    if not marker_para:
        return

    ref = marker_para._p  # 起始插入点

    for section in (sections or []):
        title = section.get('title', '').strip()
        if title:
            ref = _insert_heading_content(doc, ref, title, level=1)

        for block in section.get('blocks', []):
            btype = block.get('type', '')
            if btype == 'heading':
                ref = _insert_heading_content(
                    doc, ref,
                    block.get('text', ''),
                    level=block.get('level', 2)
                )
            elif btype == 'text':
                html = block.get('html', '')
                if html and html.strip():
                    ref = html_to_docx(html, doc, ref)
            elif btype == 'table':
                rows_data = block.get('data') or []
                if rows_data:
                    tbl_el, tail = _insert_table(
                        doc, ref, rows_data,
                        header=block.get('header', False)
                    )
                    ref = tail
            elif btype == 'image':
                img_path = (block.get('path') or '').strip()
                if img_path:
                    ref = _insert_image(
                        doc, ref, img_path,
                        width_cm=block.get('width', 0) or 0,
                        caption=block.get('caption', '')
                    )

    # 无论是否有内容，都删除 marker 段落
    marker_para._p.getparent().remove(marker_para._p)


def _insert_heading_content(doc, ref, text, level=1):
    """插入标题段落，返回新的 ref 元素"""
    new_p, new_ref = _insert_paragraph(doc, ref)
    style_name = f'Heading {level}'
    try:
        new_p.style = doc.styles[style_name]
    except (KeyError, IndexError):
        _apply_manual_heading_format(new_p, level)
    run = new_p.add_run(text)
    run.bold = True
    if level == 1 and not _has_heading_style(doc, 1):
        run.font.size = Pt(18)
    return new_ref


def _has_heading_style(doc, level):
    """检查文档是否有指定级别的标题样式"""
    try:
        _ = doc.styles[f'Heading {level}']
        return True
    except (KeyError, IndexError):
        return False


def has_body_marker(docx_path):
    """
    检查模板文件是否包含 {{__BODY__}} 标记。

    返回 bool。
    """
    from docxtpl import DocxTemplate
    try:
        doc = DocxTemplate(docx_path)
        undeclared = doc.get_undeclared_template_variables()
        return '__BODY__' in undeclared
    except Exception:
        # 回退：直接搜索 XML
        import zipfile
        import re
        with zipfile.ZipFile(docx_path, 'r') as z:
            for name in z.namelist():
                if name.startswith('word/') and name.endswith('.xml'):
                    content = z.read(name).decode('utf-8', errors='ignore')
                    if '__BODY__' in content:
                        return True
        return False
