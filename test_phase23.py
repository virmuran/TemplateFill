# -*- coding: utf-8 -*-
"""Phase 2/3 端到端测试：表格块 + 图片块 + HTML 表格 + 缺失图片兜底"""
import os
import sys
import zlib
import struct
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))


def make_png(path, w=120, h=80):
    """手工构造一张渐变 PNG（不依赖 PIL）"""
    def chunk(tag, data):
        c = struct.pack('>I', len(data)) + tag + data
        return c + struct.pack('>I', zlib.crc32(tag + data) & 0xffffffff)

    raw = b''
    for y in range(h):
        raw += b'\x00'  # filter type 0
        for x in range(w):
            raw += bytes(((x * 255) // w, (y * 255) // h, 180))

    ihdr = struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)
    with open(path, 'wb') as f:
        f.write(b'\x89PNG\r\n\x1a\n')
        f.write(chunk(b'IHDR', ihdr))
        f.write(chunk(b'IDAT', zlib.compress(raw)))
        f.write(chunk(b'IEND', b''))


def main():
    tmpdir = tempfile.mkdtemp(prefix='tf_p23_')
    img_path = os.path.join(tmpdir, 'test_img.png')
    make_png(img_path, w=600, h=400)  # 大图，验证 15cm 自动上限

    tpl_path = os.path.join(tmpdir, 'tpl.docx')
    out_path = os.path.join(tmpdir, 'out.docx')

    from doc_filler import create_sample_template, fill_template, has_body_marker
    create_sample_template(tpl_path)
    assert has_body_marker(tpl_path), '模板应含 {{__BODY__}}'

    sections = [
        {
            'title': '第一章 项目概况',
            'blocks': [
                {'type': 'heading', 'level': 2, 'text': '1.1 基本情况'},
                {'type': 'text', 'html': '<p>本项目为 <b>测试项目</b>。</p>'},
                {'type': 'table', 'header': True,
                 'data': [['项目', '数值', '单位'],
                          ['产能', '10000', 't/y'],
                          ['收率', '94.1', '%']]},
                {'type': 'text', 'html': '<p>上表为关键参数。</p>'},
                {'type': 'image', 'path': img_path, 'width': 0,
                 'caption': '图 1 测试渐变图'},
                {'type': 'text', 'html': '<p>图片之后的内容。</p>'},
            ]
        },
        {
            'title': '第二章 指定宽度图片',
            'blocks': [
                {'type': 'image', 'path': img_path, 'width': 8.0, 'caption': ''},
            ]
        },
        {
            'title': '第三章 缺失图片兜底',
            'blocks': [
                {'type': 'image', 'path': os.path.join(tmpdir, 'nope.png'),
                 'width': 10, 'caption': ''},
            ]
        },
        {
            'title': '第四章 HTML 表格（富文本粘贴）',
            'blocks': [
                {'type': 'text',
                 'html': '<table><tr><th>列A</th><th>列B</th></tr>'
                         '<tr><td>1</td><td>2</td></tr></table>'},
            ]
        },
    ]

    ctx = {t: f'[v]' for t in
           ['title', 'department', 'date', 'author', 'reviewer',
            'project_name', 'doc_number', 'overview', 'technical_plan',
            'company', 'phone']}
    fill_template(tpl_path, ctx, output_path=out_path, sections=sections)
    print('生成:', out_path)

    # ── 验证 ──
    from docx import Document
    from docx.shared import Cm
    doc = Document(out_path)

    results = []

    # 1) 表格数量：模板自带 1（基本信息）+ 章节表格 1 + HTML 表格 1 = 3
    n_tables = len(doc.tables)
    results.append(('表格总数 == 3', n_tables == 3, f'实际 {n_tables}'))

    # 2) 章节表格内容与表头
    sec_table = doc.tables[1]
    r0 = [c.text for c in sec_table.rows[0].cells]
    r1 = [c.text for c in sec_table.rows[1].cells]
    results.append(('表格 3行x3列',
                    len(sec_table.rows) == 3 and len(sec_table.columns) == 3,
                    f'{len(sec_table.rows)}x{len(sec_table.columns)}'))
    results.append(('表头行内容', r0 == ['项目', '数值', '单位'], str(r0)))
    results.append(('数据行内容', r1 == ['产能', '10000', 't/y'], str(r1)))
    head_runs = sec_table.rows[0].cells[0].paragraphs[0].runs
    results.append(('表头加粗', bool(head_runs and head_runs[0].bold), ''))

    # 3) HTML 表格
    html_table = doc.tables[2]
    hr0 = [c.text for c in html_table.rows[0].cells]
    results.append(('HTML表格 2x2',
                    len(html_table.rows) == 2 and hr0 == ['列A', '列B'],
                    f'{len(html_table.rows)}行, 首行{hr0}'))

    # 4) 图片：2 张有效（600px@96dpi≈15.9cm → 截到 15cm；指定 8cm）
    shapes = doc.inline_shapes
    results.append(('内嵌图片数 == 2', len(shapes) == 2, f'实际 {len(shapes)}'))
    if len(shapes) >= 2:
        s0, s1 = shapes[0], shapes[1]
        results.append(('大图截宽 15cm', abs(s0.width - Cm(15)) < Cm(0.2),
                        f'{s0.width.cm:.2f}cm'))
        results.append(('比例保持', abs(s0.height / s0.width - 400 / 600) < 0.02,
                        f'h/w={s0.height / s0.width:.3f}'))
        results.append(('指定宽度 8cm', abs(s1.width - Cm(8)) < Cm(0.2),
                        f'{s1.width.cm:.2f}cm'))

    # 5) 图注 / 缺失占位 / 标记清除
    all_text = '\n'.join(p.text for p in doc.paragraphs)
    results.append(('图注存在', '图 1 测试渐变图' in all_text, ''))
    results.append(('缺失占位', '图片缺失' in all_text, ''))
    results.append(('BODY标记已删', 'BODY_CONTENT_MARKER' not in all_text, ''))
    results.append(('章节标题H1', '第三章 缺失图片兜底' in all_text, ''))

    print()
    ok = True
    for name, passed, detail in results:
        mark = 'PASS' if passed else 'FAIL'
        if not passed:
            ok = False
        print(f'[{mark}] {name}' + (f'  ({detail})' if detail else ''))

    print()
    print('段落结构预览:')
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            print(f'  [{i}] {p.style.name}: {p.text[:40]}')

    print()
    print('总体:', 'ALL PASS' if ok else 'HAS FAILURES')
    return 0 if ok else 1


if __name__ == '__main__':
    sys.exit(main())
